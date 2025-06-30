import os
import tempfile
from typing import Tuple, Optional
import PyPDF2
import docx
from fastapi import UploadFile
from pptx import Presentation
import logging
from PIL import Image
import pytesseract
from docx import Document
import pandas as pd
from io import BytesIO
import chardet
from app.configuration.files import MAX_FILE_SIZE, ALLOWED_FILE_TYPES
from app.utils.benchmark import benchmark

logger = logging.getLogger(__name__)

BINARY_CONTENT_TYPES = {
    'application/pdf',
    'application/msword',
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'application/vnd.ms-excel',
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    'application/vnd.ms-powerpoint'
}

@benchmark.time_it('text')
async def process_uploaded_file(file: UploadFile, file_type: str) -> dict:
    if file_type not in ALLOWED_FILE_TYPES:
        raise ValueError(f"Invalid file type: {file_type}")
    
    if file.content_type not in ALLOWED_FILE_TYPES[file_type]:
        raise ValueError(f"Unsupported content type: {file.content_type} for type {file_type}")

    file.file.seek(0, 2)
    file_size = file.file.tell()
    file.file.seek(0)
    
    if file_size > MAX_FILE_SIZE:
        raise ValueError(f"File size {file_size} exceeds maximum allowed size {MAX_FILE_SIZE}")

    try:
        content_bytes = await file.read()
        content = ""  # По умолчанию пустое содержимое
        
        # Для изображений возвращаем base64
        if file_type == 'image':
            import base64
            content = f"data:{file.content_type};base64,{base64.b64encode(content_bytes).decode('utf-8')}"
            return {
                "type": file_type,
                "name": file.filename,
                "content": content,
                "content_type": file.content_type,
                "size": file_size,
                "image_data": content_bytes
            }
        
        # Для текстовых файлов определяем кодировку
        if file.content_type in ALLOWED_FILE_TYPES:
            detected = chardet.detect(content_bytes)
            encoding = detected['encoding'] or 'utf-8'
            try:
                content = content_bytes.decode(encoding, errors='replace')
            except Exception:
                try:
                    content = content_bytes.decode('utf-8', errors='replace')
                except Exception:
                    content = content_bytes.decode('latin-1', errors='replace')
        
        # Для бинарных форматов извлекаем текст
        elif file.content_type in BINARY_CONTENT_TYPES:
            if file.content_type in [
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                'application/msword'
            ]:
                content = extract_text_from_docx_bytes(content_bytes)
            elif file.content_type == 'application/pdf':
                content = extract_text_from_pdf_bytes(content_bytes)
            elif file.content_type in [
                'application/vnd.openxmlformats-officedocument.presentationml.presentation',
                'application/vnd.ms-powerpoint'
            ]:
                content = extract_text_from_pptx_bytes(content_bytes)
        text_metrics = {}
        if content:
            from app.utils.text_metrics import calculate_text_metrics
            text_metrics = calculate_text_metrics(content)
        return {
            "type": file_type,
            "name": file.filename,
            "content": content,
            "content_type": file.content_type,
            "size": file_size,
            "metrics": text_metrics
        }
    except Exception as e:
        raise ValueError(f"Error processing file: {str(e)}")
    finally:
        await file.close()
@benchmark.time_it('pdf')
def extract_text_from_pdf_bytes(content: bytes) -> str:
    """Извлекает текст из PDF из байтов"""
    text = ""
    try:
        with BytesIO(content) as bytes_io:
            reader = PyPDF2.PdfReader(bytes_io)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        logger.error(f"PDF extraction error: {str(e)}", exc_info=True)
    return text

@benchmark.time_it('docx')
def extract_text_from_docx_bytes(content: bytes) -> str:
    """Извлекает текст из DOCX файла из байтов"""
    try:
        with BytesIO(content) as bytes_io:
            doc = docx.Document(bytes_io)
            return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        logger.error(f"DOCX extraction error: {str(e)}", exc_info=True)
        return ""
    
@benchmark.time_it('pptx')
def extract_text_from_pptx_bytes(content: bytes) -> str:
    """Извлекает текст из PPTX из байтов"""
    try:
        with BytesIO(content) as bytes_io:
            prs = Presentation(bytes_io)
            text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
            return "\n".join(text)
    except Exception as e:
        logger.error(f"PPTX extraction error: {str(e)}", exc_info=True)
        return ""

def extract_text_from_file(file_path: str, file_type: str) -> Tuple[Optional[str], Optional[bytes]]:
    """Извлекает текст из файлов различных форматов"""
    try:
        text = None
        image_data = None
        
        if file_type == 'application/pdf':
            text = extract_text_from_pdf(file_path)
        elif file_type == 'text/plain':
            with open(file_path, 'rb') as f:
                content_bytes = f.read()
                detected = chardet.detect(content_bytes)
                encoding = detected['encoding'] or 'utf-8'
                try:
                    return content_bytes.decode(encoding, errors='replace'), None
                except Exception:
                    try:
                        return content_bytes.decode('utf-8', errors='replace'), None
                    except Exception:
                        return content_bytes.decode('latin-1', errors='replace'), None
        elif file_type in [
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'application/msword'
        ]:
            text = extract_text_from_docx(file_path)
        elif file_type in [
            'application/vnd.openxmlformats-officedocument.presentationml.presentation',
            'application/vnd.ms-powerpoint'
        ]:
            text = extract_text_from_pptx(file_path)
        elif file_type.startswith('image/'):
            text = extract_text_from_image(file_path)
            with open(file_path, 'rb') as f:
                image_data = f.read()
        
        return text, image_data
    except Exception as e:
        logger.error(f"Error processing file {file_path}: {str(e)}", exc_info=True)
        return None, None

def extract_text_from_pdf(file_path: str) -> str:
    """Извлекает текст из PDF"""
    text = ""
    try:
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        logger.error(f"PDF extraction error: {str(e)}", exc_info=True)
    return text

def extract_text_from_docx(file_path: str) -> str:
    """Извлекает текст из DOCX"""
    try:
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        logger.error(f"DOCX extraction error: {str(e)}", exc_info=True)
        # Fallback для поврежденных DOCX файлов
        try:
            with open(file_path, 'rb') as f:
                content = f.read()
                detected = chardet.detect(content)
                encoding = detected['encoding'] or 'utf-8'
                return content.decode(encoding, errors='replace')
        except Exception as fallback_e:
            logger.error(f"DOCX fallback extraction failed: {str(fallback_e)}", exc_info=True)
            return ""

def extract_text_from_pptx(file_path: str) -> str:
    """Извлекает текст из PPTX"""
    try:
        prs = Presentation(file_path)
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
        return "\n".join(text)
    except Exception as e:
        logger.error(f"PPTX extraction error: {str(e)}", exc_info=True)
        return ""
@benchmark.time_it('image')
def extract_text_from_image(file_path: str) -> str:
    """Извлекает текст из изображения с помощью OCR"""
    try:
        img = Image.open(file_path)
        return pytesseract.image_to_string(img)
    except Exception as e:
        logger.error(f"OCR failed: {str(e)}", exc_info=True)
        return ""